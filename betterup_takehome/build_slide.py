from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

BLUE = RGBColor(0x4A, 0x69, 0xBD)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG_LIGHT = RGBColor(0xF5, 0xF6, 0xFA)
ACCENT_BG = RGBColor(0xE8, 0xED, 0xF7)


def add_shape(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.fill.solid()
        shape.line.fore_color.rgb = line_color
        shape.line.width = Pt(1)
    return shape


def add_text_box(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def add_run(paragraph, text, font_size=10, bold=False, color=DARK, font_name="Calibri"):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return run


# ── Top accent bar ──
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), fill_color=BLUE)

# ── Title ──
title_box = add_text_box(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7))
p = title_box.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
add_run(p, "TechCorp's Coaching Investment Is Delivering", font_size=26, bold=True, color=DARK)
add_run(p, " — Recommend Renewal", font_size=26, bold=False, color=BLUE)

# ── Subtitle ──
sub_box = add_text_box(slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.35))
p = sub_box.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
add_run(p, "Mid-Year Business Value Review  |  July 2025  |  500 Licenses  |  6 Months In", font_size=12, color=GRAY)

# ── Thin divider under subtitle ──
add_shape(slide, Inches(0.6), Inches(1.35), Inches(12.1), Inches(0.015), fill_color=RGBColor(0xDD, 0xDD, 0xDD))

# ════════════════════════════════════════════════════
# LEFT COLUMN: Is it being used?
# ════════════════════════════════════════════════════

left_x = Inches(0.6)
col_top = Inches(1.6)
col_w = Inches(5.6)

# Section header
hdr = add_text_box(slide, left_x, col_top, col_w, Inches(0.35))
p = hdr.text_frame.paragraphs[0]
add_run(p, "IS THE PROGRAM BEING USED?", font_size=11, bold=True, color=BLUE)

# ── Stat 1: 298 active users ──
y = Inches(2.05)
s1 = add_text_box(slide, left_x, y, col_w, Inches(0.55))
p = s1.text_frame.paragraphs[0]
add_run(p, "298", font_size=36, bold=True, color=DARK)
add_run(p, "  active coaching users", font_size=14, color=GRAY)
p2 = s1.text_frame.add_paragraph()
add_run(p2, "60% of purchased licenses actively engaged in coaching sessions", font_size=10.5, color=GRAY)

# ── Stat 2: 4.2 avg sessions ──
y = Inches(2.85)
s2 = add_text_box(slide, left_x, y, col_w, Inches(0.55))
p = s2.text_frame.paragraphs[0]
add_run(p, "4.2", font_size=36, bold=True, color=DARK)
add_run(p, "  avg sessions per user", font_size=14, color=GRAY)
p2 = s2.text_frame.add_paragraph()
add_run(p2, "Sustained, repeated engagement — not one-and-done", font_size=10.5, color=GRAY)

# ── Stat 3: Projection ──
y = Inches(3.65)
proj_bg = add_shape(slide, Inches(0.4), y, Inches(5.9), Inches(0.85), fill_color=ACCENT_BG)
proj_bg.line.fill.background()

s3 = add_text_box(slide, left_x, Inches(3.72), Inches(5.6), Inches(0.75))
p = s3.text_frame.paragraphs[0]
add_run(p, "~75%", font_size=28, bold=True, color=BLUE)
add_run(p, "  projected utilization by renewal", font_size=13, color=DARK)
p2 = s3.text_frame.add_paragraph()
add_run(p2, "On current momentum alone, without deploying new engagement strategies", font_size=10.5, color=GRAY)

# ════════════════════════════════════════════════════
# RIGHT COLUMN: Are employees finding value?
# ════════════════════════════════════════════════════

right_x = Inches(7.1)
right_w = Inches(5.6)

# Vertical divider
add_shape(slide, Inches(6.65), Inches(1.6), Inches(0.015), Inches(3.2), fill_color=RGBColor(0xDD, 0xDD, 0xDD))

# Section header
hdr2 = add_text_box(slide, right_x, col_top, right_w, Inches(0.35))
p = hdr2.text_frame.paragraphs[0]
add_run(p, "ARE EMPLOYEES FINDING VALUE?", font_size=11, bold=True, color=BLUE)

# ── Stat: NPS +45 ──
y = Inches(2.05)
n1 = add_text_box(slide, right_x, y, right_w, Inches(0.55))
p = n1.text_frame.paragraphs[0]
add_run(p, "NPS +45", font_size=36, bold=True, color=DARK)
p2 = n1.text_frame.add_paragraph()
add_run(p2, "Top quartile — industry average is ~+30", font_size=10.5, color=GRAY)

# ── Stat: Satisfaction ──
y = Inches(2.85)
n2 = add_text_box(slide, right_x, y, right_w, Inches(0.7))
p = n2.text_frame.paragraphs[0]
add_run(p, "4.3", font_size=36, bold=True, color=DARK)
add_run(p, " / 5.0", font_size=18, color=GRAY)
add_run(p, "  overall satisfaction", font_size=14, color=GRAY)
p2 = n2.text_frame.add_paragraph()
add_run(p2, "4.5 / 5.0", font_size=12, bold=True, color=DARK)
add_run(p2, "  \"My coach helps me achieve my goals\"", font_size=10.5, color=GRAY)

# ── Stat: Survey response rate ──
y = Inches(3.72)
resp_bg = add_shape(slide, Inches(6.9), Inches(3.65), Inches(5.9), Inches(0.85), fill_color=ACCENT_BG)
resp_bg.line.fill.background()

n3 = add_text_box(slide, right_x, y, right_w, Inches(0.75))
p = n3.text_frame.paragraphs[0]
add_run(p, "52%", font_size=28, bold=True, color=BLUE)
add_run(p, "  survey response rate", font_size=13, color=DARK)
p2 = n3.text_frame.add_paragraph()
add_run(p2, "Well above the typical 25–35% — employees are engaged enough to give feedback", font_size=10.5, color=GRAY)

# ════════════════════════════════════════════════════
# BOTTOM: Recommendation
# ════════════════════════════════════════════════════

rec_top = Inches(5.1)
rec_bg = add_shape(slide, Inches(0), rec_top, Inches(13.333), Inches(1.6), fill_color=RGBColor(0x2C, 0x3E, 0x6B))

# Recommendation header
rec_hdr = add_text_box(slide, Inches(0.6), Inches(5.25), Inches(12), Inches(0.35))
p = rec_hdr.text_frame.paragraphs[0]
add_run(p, "RECOMMENDATION", font_size=11, bold=True, color=RGBColor(0xA0, 0xB4, 0xE0))

# Recommendation main text
rec_main = add_text_box(slide, Inches(0.6), Inches(5.6), Inches(11.5), Inches(0.5))
p = rec_main.text_frame.paragraphs[0]
add_run(p, "Renew.", font_size=22, bold=True, color=WHITE)
add_run(p, "  The foundation is strong — high satisfaction, deep engagement, and sustained adoption growth.",
        font_size=15, color=RGBColor(0xD0, 0xD8, 0xEE))

# Recommendation supporting text
rec_sub = add_text_box(slide, Inches(0.6), Inches(6.15), Inches(11.5), Inches(0.4))
p = rec_sub.text_frame.paragraphs[0]
add_run(p, "A targeted activation strategy, starting now, can strengthen the growth trend and further maximize ROI heading into renewal.",
        font_size=12, color=RGBColor(0xA0, 0xB4, 0xE0))

# ── Footer ──
footer = add_text_box(slide, Inches(0.6), Inches(7.05), Inches(12), Inches(0.3))
p = footer.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
add_run(p, "BetterUp  |  Confidential  |  Prepared for TechCorp Leadership", font_size=8, color=LIGHT_GRAY)

output_path = "/Users/truettbloxsom/Documents/Repositories/yt-examples/betterup_takehome/deliverable_1_executive_slide.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
